#!/usr/bin/env python
import os
import sys
import threading
import uuid
import time
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
os.chdir(BASE_DIR)
sys.path.insert(0, str(BASE_DIR / "src"))
sys.path.insert(0, str(BASE_DIR))

from dotenv import load_dotenv

load_dotenv(BASE_DIR / ".env")

from flask import Flask, jsonify, redirect, render_template, request, send_file, url_for
from markupsafe import Markup, escape
import io
import json
import re
from content_agent.main import run
from content_agent.helpers import keyword_slug
from content_agent.crew import default_llm

from docx import Document
from docx.shared import Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

app = Flask(__name__)
app.secret_key = "speedbots_content_web"

URL_RE = re.compile(r"(https?://[^\s<)]+)")
REL_URL_RE = re.compile(r"(/[\w]+(?:-[\w]+)+)")

HEADING_STARTS = ("SEO TITLE:", "META DESCRIPTION:", "CTA BUTTONS:", "SEO REPORT")


def is_heading_line(s):
    if not s or len(s) < 6:
        return False
    letters = [ch for ch in s if ch.isalpha()]
    if not letters:
        return False
    return s == s.upper()


def _docx_hyperlink(paragraph, text, url):
    """Append a clickable hyperlink to *paragraph*."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1d4ed8")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run_el.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run_el.append(t)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _add_text_with_links(paragraph, text):
    """Split *text* on URLs and markdown **bold**, adding runs/hyperlinks accordingly."""
    parts = URL_RE.split(text)
    for part in parts:
        if URL_RE.fullmatch(part):
            _docx_hyperlink(paragraph, part, part)
            continue
        pos = 0
        for m in re.finditer(r"\*\*(.+?)\*\*", part):
            if m.start() > pos:
                paragraph.add_run(part[pos:m.start()])
            r = paragraph.add_run(m.group(1))
            r.bold = True
            pos = m.end()
        if pos < len(part):
            paragraph.add_run(part[pos:])


def _blog_to_docx(content, doc_title="Speedbots | Blog"):
    """Convert the raw blog text into a python-docx Document."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if doc_title:
        p = doc.add_paragraph(doc_title)
        p.style = doc.styles["Title"]
        p.runs[0].font.size = Pt(20)
        p.runs[0].font.bold = True

    for raw in content.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        s = line.strip()

        # front-matter lines
        if s.startswith("SEO TITLE:"):
            continue
        if s.startswith("META DESCRIPTION:"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(s.split(":", 1)[1].strip())
            run.italic = True
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            continue
        if s.startswith("CTA BUTTONS:"):
            continue

        # headings
        if is_heading_line(s):
            doc.add_heading(s, level=2)
            continue
        if s.startswith(("SEO REPORT", "REAL RESULTS", "FREQUENTLY ASKED")):
            doc.add_heading(s, level=2)
            continue

        # bullets
        if s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_text_with_links(p, s[2:].strip())
            continue

        # numbered steps
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_text_with_links(p, m.group(2))
            continue

        # bold headings (from **...** markdown)
        if s.startswith("**") and s.endswith("**"):
            doc.add_heading(s.strip("*"), level=2)
            continue

        # normal paragraph
        p = doc.add_paragraph()
        _add_text_with_links(p, s)

    return doc


@app.template_filter("blog")
def blog_filter(text):
    """Convert a plain-text blog into HTML: highlighted clickable links, bold headings."""
    if not text:
        return Markup("")
    out = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        s = line.strip()
        esc = str(escape(s))
        anchors = []
        counter = [0]

        def stash(fmt, url=False):
            def sub(match):
                token = match.group(1)
                if url:
                    m = re.match(r"^(.*?)([.,;:!?)]*)$", token)
                    core, trailing = m.group(1), m.group(2)
                    anchors.append(fmt.format(core))
                    placeholder = f"\u0000{len(anchors) - 1}\u0000"
                    return placeholder + trailing
                anchors.append(fmt.format(token))
                placeholder = f"\u0000{len(anchors) - 1}\u0000"
                return placeholder
            return sub

        esc = URL_RE.sub(stash('<a href="{0}" target="_blank" rel="noopener" class="blog-link">{0}</a>', url=True), esc)
        esc = REL_URL_RE.sub(
            stash('<a href="https://speedbots.com{0}" target="_blank" rel="noopener" class="blog-link">{0}</a>'),
            esc,
        )
        esc = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", esc)
        for i, a in enumerate(anchors):
            esc = esc.replace(f"\u0000{i}\u0000", a)
        if is_heading_line(s) or s.startswith(HEADING_STARTS):
            esc = f"<strong>{esc}</strong>"
        out.append(esc)
    return Markup("<br>\n".join(out))

BLOGS_DIR = BASE_DIR / "output" / "blogs"
REPORTS_DIR = BASE_DIR / "output" / "reports"
IMAGES_DIR = BASE_DIR / "output" / "images"
HISTORY_FILE = BASE_DIR / "output" / "blog_history.json"

JOBS = {}
JOB_LOCK = threading.Lock()

COUNTRIES = [
    "All over the world",
    "Malaysia",
    "Singapore",
    "Indonesia",
    "Thailand",
    "Philippines",
    "Australia",
    "Coimbatore",
    "India",
]

SUBTOPICS = ["basics", "standard", "roi", "educational", "ai"]


def _warm_model():
    try:
        llm = default_llm()
        llm.call("Respond with only the word: ready")
    except Exception:
        pass


def _generate(job_id, keyword, country, internal_links, subtopics="standard"):
    try:
        JOBS[job_id]["status"] = "running"
        slug = keyword_slug(keyword)
        blog_path = BLOGS_DIR / f"{slug}.txt"
        for p in [blog_path, REPORTS_DIR / "strategy.json", REPORTS_DIR / "research.md",
                  REPORTS_DIR / "seo_report.txt", REPORTS_DIR / "quality_review.md"]:
            if p.exists():
                p.unlink()
        for ext in ("png", "jpg", "jpeg"):
            ip = IMAGES_DIR / f"{slug}.{ext}"
            if ip.exists():
                ip.unlink()

        with JOB_LOCK:
            run(keyword, country, internal_links, subtopics)

        if blog_path.exists():
            content = blog_path.read_text(encoding="utf-8")
            image_path = IMAGES_DIR / f"{slug}.png"
            if not image_path.exists():
                image_path = IMAGES_DIR / f"{slug}.jpg"
            JOBS[job_id].update(status="done", blog=content, path=str(blog_path),
                                slug=slug, words=len(content.split()),
                                image=str(image_path) if image_path.exists() else None)
            _record_history(slug, country)
        else:
            JOBS[job_id].update(status="error", error="Blog file was not created.")
    except Exception as e:
        JOBS[job_id].update(status="error", error=str(e))


def _slug_image(slug):
    for ext in ("png", "jpg", "jpeg"):
        ip = IMAGES_DIR / f"{slug}.{ext}"
        if ip.exists():
            return str(ip)
    return None


def _extract_title(text, fallback=""):
    m = re.search(r"SEO TITLE: ?(.*)", text)
    if m:
        return m.group(1).strip().lstrip(" *#*")
    return fallback


def _load_history():
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _save_history(hist):
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(hist, f, indent=2, sort_keys=True)


def _seed_history():
    items = {}
    for p in BLOGS_DIR.glob("*.txt"):
        try:
            text = p.read_text(encoding="utf-8")
        except Exception:
            continue
        if "SEO TITLE:" not in text:
            continue
        if not _slug_image(p.stem):
            continue
        items[p.stem] = {
            "slug": p.stem,
            "title": _extract_title(text, p.stem.replace("-", " ").title()),
            "words": len(text.split()),
            "mtime": p.stat().st_mtime,
            "country": "",
        }
    _save_history(items)


def _record_history(slug, country=""):
    p = BLOGS_DIR / f"{slug}.txt"
    if not p.exists():
        return
    text = p.read_text(encoding="utf-8")
    if "SEO TITLE:" not in text:
        return
    hist = _load_history()
    hist[slug] = {
        "slug": slug,
        "title": _extract_title(text, slug.replace("-", " ").title()),
        "words": len(text.split()),
        "mtime": time.time(),
        "country": country,
    }
    _save_history(hist)


def _blog_list():
    hist = _load_history()
    if not hist:
        _seed_history()
        hist = _load_history()
    items = []
    for slug in sorted(hist, key=lambda s: hist[s].get("mtime", 0), reverse=True):
        p = BLOGS_DIR / f"{slug}.txt"
        if not p.exists():
            continue
        try:
            text = p.read_text(encoding="utf-8")
        except Exception:
            continue
        if "SEO TITLE:" not in text:
            continue
        e = hist[slug]
        items.append({
            "slug": slug,
            "title": _extract_title(text, e.get("title", slug.replace("-", " ").title())),
            "words": len(text.split()),
            "mtime": time.strftime("%b %d, %Y %H:%M", time.localtime(e.get("mtime", p.stat().st_mtime))),
            "image": _slug_image(slug),
        })
    return items


@app.route("/")
def index():
    return render_template("index.html", countries=COUNTRIES)


@app.route("/history")
def history():
    return render_template("history.html", blogs=_blog_list())


@app.route("/view/<slug>")
def view(slug):
    p = BLOGS_DIR / f"{slug}.txt"
    if not p.exists():
        return redirect(url_for("history"))
    text = p.read_text(encoding="utf-8")
    if "SEO TITLE:" not in text:
        return redirect(url_for("history"))
    title = _extract_title(text, slug.replace("-", " ").title())
    job = {"status": "done", "slug": slug, "title": title, "blog": text,
           "words": len(text.split()), "image": _slug_image(slug), "keyword": title, "country": ""}
    return render_template("result.html", job=job, job_id=slug)


@app.route("/generate", methods=["POST"])
def generate():
    keyword = request.form.get("keyword", "").strip()
    country = request.form.get("country", "Malaysia").strip()
    internal_links = request.form.get("internal_links", "").strip()
    subtopics = request.form.get("subtopics", "").strip() or "basics"
    if not keyword:
        return redirect(url_for("index"))

    job_id = uuid.uuid4().hex[:10]
    JOBS[job_id] = {"status": "queued", "keyword": keyword, "country": country, "subtopics": subtopics}
    thread = threading.Thread(target=_generate, args=(job_id, keyword, country, internal_links, subtopics), daemon=True)
    thread.start()
    return redirect(url_for("status", job_id=job_id))


@app.route("/status/<job_id>")
def status(job_id):
    if job_id not in JOBS:
        return redirect(url_for("index"))
    return render_template("status.html", job_id=job_id, job=JOBS[job_id])


@app.route("/api/status/<job_id>")
def api_status(job_id):
    if job_id not in JOBS:
        return jsonify({"error": "unknown job"}), 404
    return jsonify(JOBS[job_id])


@app.route("/result/<job_id>")
def result(job_id):
    if job_id not in JOBS:
        return redirect(url_for("index"))
    job = JOBS[job_id]
    if job["status"] != "done":
        return redirect(url_for("status", job_id=job_id))
    return render_template("result.html", job=job, job_id=job_id)


@app.route("/download/<job_id>")
def download(job_id):
    content = slug = None
    if job_id in JOBS and JOBS[job_id]["status"] == "done":
        job = JOBS[job_id]
        content = job.get("blog") or ""
        slug = job.get("slug") or "blog"
    elif (BLOGS_DIR / f"{job_id}.txt").exists():
        text = (BLOGS_DIR / f"{job_id}.txt").read_text(encoding="utf-8")
        if "SEO TITLE:" in text:
            content = text
            slug = job_id
    if content is None:
        return redirect(url_for("index"))
    doc_title = "Speedbots | Blog"
    et = _extract_title(content)
    if et:
        doc_title = et
    doc = _blog_to_docx(content, doc_title)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{slug}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.route("/image/<job_id>")
def image(job_id):
    path = None
    if job_id in JOBS and JOBS[job_id].get("image") and Path(JOBS[job_id]["image"]).exists():
        path = JOBS[job_id]["image"]
    else:
        path = _slug_image(job_id)
    if not path:
        return redirect(url_for("index"))
    p = Path(path)
    if p.suffix.lower() in (".jpg", ".jpeg"):
        return send_file(p, mimetype="image/jpeg")
    return send_file(p, mimetype="image/png")


if __name__ == "__main__":
    os.environ.setdefault("FLASK_ENV", "development")
    threading.Thread(target=_warm_model, daemon=True).start()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5001")), debug=False, threaded=True)